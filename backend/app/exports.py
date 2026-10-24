import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_excel_report(transactions, loans, analytics) -> io.BytesIO:
    """
    Generates a clean, simple, unbloated Excel file ready for formulas & sorting.
    """
    wb = Workbook()
    
    # Header styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    # Sheet 1: Tranzaksiyalar
    ws1 = wb.active
    ws1.title = "Tranzaksiyalar"
    
    headers1 = ["ID", "Sana va Vaqt", "Turi", "Kategoriya", "Summa (so'm)", "Izoh / Ifoda"]
    ws1.append(headers1)
    
    for col_idx, col_name in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    
    type_labels = {
        "income": "Daromad (+)",
        "expense": "Xarajat (-)",
        "loan_payment": "Kredit to'lovi",
        "unknown_loss": "Qora tuynuk (Yo'qolgan)",
        "goal_deposit": "Jamg'arma"
    }
    
    for row_idx, t in enumerate(transactions, 2):
        created_str = t.created_at.strftime("%Y-%m-%d %H:%M") if t.created_at else ""
        row = [
            t.id,
            created_str,
            type_labels.get(t.type, t.type),
            t.category or "Boshqa",
            t.amount,
            t.note or t.raw_expression or ""
        ]
        ws1.append(row)
        
        # Format amount cell as currency
        amount_cell = ws1.cell(row=row_idx, column=5)
        amount_cell.number_format = '#,##0'
        
        for c in range(1, 7):
            ws1.cell(row=row_idx, column=c).border = thin_border
            
    # Auto-adjust column widths
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    # Sheet 2: Kreditlar
    ws2 = wb.create_sheet(title="Kreditlar va Qarzlar")
    headers2 = ["Kredit Nomi", "Jami Summa (so'm)", "Oylik To'lov", "Qolgan Qarz", "To'lov Kuni", "Holati"]
    ws2.append(headers2)
    
    for col_idx, col_name in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        
    for row_idx, l in enumerate(loans, 2):
        ws2.append([
            l.title,
            l.total_amount,
            l.monthly_payment,
            l.remaining_amount,
            f"Har oyning {l.due_day}-sanasi",
            "Faol" if l.is_active else "To'langan"
        ])
        ws2.cell(row=row_idx, column=2).number_format = '#,##0'
        ws2.cell(row=row_idx, column=3).number_format = '#,##0'
        ws2.cell(row=row_idx, column=4).number_format = '#,##0'
        for c in range(1, 7):
            ws2.cell(row=row_idx, column=c).border = thin_border
            
    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)
        
    # Sheet 3: Oylik Xulosa
    ws3 = wb.create_sheet(title="Oylik Xulosa")
    ws3.append(["Ko'rsatkich", "Summa / Qiymat"])
    ws3.cell(row=1, column=1).font = header_font
    ws3.cell(row=1, column=1).fill = header_fill
    ws3.cell(row=1, column=2).font = header_font
    ws3.cell(row=1, column=2).fill = header_fill
    
    summary_data = [
        ("Oy nomi", analytics.month_name),
        ("Jami Daromad", analytics.total_income),
        ("Jami Xarajatlar", analytics.total_expense),
        ("Kreditlarga To'langan", analytics.total_loans_paid),
        ("Qora tuynuk (Noma'lum yo'qolgan)", analytics.mystery_loss_total),
        ("Mayda-chuyda xarajatlar", analytics.small_expenses_total),
        ("Sarflangan mehnat soati", f"{analytics.work_hours_spent} soat"),
        ("Sof qoldiq (Daromad - Xarajat)", analytics.total_income - analytics.total_expense)
    ]
    
    for r_idx, (k, v) in enumerate(summary_data, 2):
        ws3.append([k, v])
        if isinstance(v, (int, float)):
            ws3.cell(row=r_idx, column=2).number_format = '#,##0'
        for c in range(1, 3):
            ws3.cell(row=r_idx, column=c).border = thin_border
            
    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws3.column_dimensions[col_letter].width = max(max_len + 4, 20)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_pdf_report(transactions, loans, analytics) -> io.BytesIO:
    """
    Generates a stunning, executive-styled financial statement in PDF.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0F172A')
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#64748B')
    )
    card_title_style = ParagraphStyle(
        'CardTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor('#64748B')
    )
    card_val_green = ParagraphStyle(
        'CardValGreen',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=16,
        textColor=colors.HexColor('#16A34A')
    )
    card_val_red = ParagraphStyle(
        'CardValRed',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=16,
        textColor=colors.HexColor('#DC2626')
    )
    card_val_blue = ParagraphStyle(
        'CardValBlue',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=16,
        textColor=colors.HexColor('#2563EB')
    )
    sec_title = ParagraphStyle(
        'SectionTitle',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1E293B')
    )
    cell_text = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#334155')
    )
    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#0F172A')
    )
    
    story = []
    
    # 1. Header Banner
    story.append(Paragraph("MAOSH VA XARAJATLAR", title_style))
    story.append(Paragraph(f"Oylik Moliyaviy Hisobot &bull; <b>{analytics.month_name}</b> &bull; Hujjat sanasi: {datetime.now().strftime('%d.%m.%Y %H:%M')}", subtitle_style))
    story.append(Spacer(1, 14))
    
    # 2. Executive Summary Cards (3 Columns)
    net_flow = analytics.total_income - analytics.total_expense
    flow_style = card_val_green if net_flow >= 0 else card_val_red
    
    cards_data = [[
        [Paragraph("JAMI DAROMAD", card_title_style), Spacer(1, 4), Paragraph(f"+{analytics.total_income:,.0f} so'm", card_val_green)],
        [Paragraph("JAMI XARAJAT", card_title_style), Spacer(1, 4), Paragraph(f"-{analytics.total_expense:,.0f} so'm", card_val_red)],
        [Paragraph("SOF QOLDIQ", card_title_style), Spacer(1, 4), Paragraph(f"{net_flow:+,.0f} so'm", flow_style)]
    ]]
    
    cards_table = Table(cards_data, colWidths=[174, 174, 174])
    cards_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,0), colors.HexColor('#F0FDF4')),
        ('BACKGROUND', (1,0), (1,0), colors.HexColor('#FEF2F2')),
        ('BACKGROUND', (2,0), (2,0), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (0,0), 1, colors.HexColor('#DCFCE7')),
        ('BOX', (1,0), (1,0), 1, colors.HexColor('#FEE2E2')),
        ('BOX', (2,0), (2,0), 1, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
    ]))
    story.append(cards_table)
    story.append(Spacer(1, 14))
    
    # 3. Mystery Loss / Black Hole Highlight Box
    loss_msg = (
        f"<b>Diqqat:</b> Ushbu oyda <b>{analytics.mystery_loss_total:,.0f} so'm</b> pulingiz qayerga ketgani noma'lum yo'qolgan! "
        f"Shuningdek, mayda-chuyda xarajatlar jami <b>{analytics.small_expenses_total:,.0f} so'm</b>ni tashkil qildi. "
        f"Bu sarflangan barcha pullar taxminan <b>{analytics.work_hours_spent} soatlik</b> mehnatingizga teng."
    )
    loss_data = [[
        Paragraph("🕳 QORA TUYNUQ VA MAYDA XARAJATLAR TAHLILI", ParagraphStyle('AlertTitle', fontName='Helvetica-Bold', fontSize=10, textColor=colors.HexColor('#9A3412'))),
    ], [
        Paragraph(loss_msg, ParagraphStyle('AlertMsg', fontName='Helvetica', fontSize=8.5, leading=12, textColor=colors.HexColor('#7C2D12')))
    ]]
    loss_table = Table(loss_data, colWidths=[522])
    loss_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#FFF7ED')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#FDBA74')),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
    ]))
    story.append(loss_table)
    story.append(Spacer(1, 14))
    
    # 4. Active Loans Status Table
    if loans:
        story.append(Paragraph("Kreditlar va Muddatli To'lovlar Holati", sec_title))
        story.append(Spacer(1, 6))
        loans_table_data = [[
            Paragraph("Kredit Nomi", cell_bold),
            Paragraph("Jami Qarz", cell_bold),
            Paragraph("Oylik To'lov", cell_bold),
            Paragraph("Qolgan Qarz", cell_bold),
            Paragraph("To'lov Sanasi", cell_bold)
        ]]
        for l in loans:
            loans_table_data.append([
                Paragraph(l.title, cell_text),
                Paragraph(f"{l.total_amount:,.0f}", cell_text),
                Paragraph(f"{l.monthly_payment:,.0f}", cell_text),
                Paragraph(f"<b>{l.remaining_amount:,.0f}</b>", cell_text),
                Paragraph(f"Har oy {l.due_day}-sana", cell_text),
            ])
        lt = Table(loans_table_data, colWidths=[150, 95, 95, 95, 87])
        lt.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
            ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ]))
        story.append(lt)
        story.append(Spacer(1, 14))
        
    # 5. Recent Transactions Sample
    story.append(Paragraph("So'nggi Tranzaksiyalar Ko'chirmasi", sec_title))
    story.append(Spacer(1, 6))
    
    tx_table_data = [[
        Paragraph("Sana", cell_bold),
        Paragraph("Turi", cell_bold),
        Paragraph("Kategoriya", cell_bold),
        Paragraph("Summa (so'm)", cell_bold),
        Paragraph("Izoh", cell_bold),
    ]]
    
    type_labels = {
        "income": "Daromad",
        "expense": "Xarajat",
        "loan_payment": "Kredit to'lovi",
        "unknown_loss": "Qora tuynuk",
        "goal_deposit": "Jamg'arma"
    }
    
    for t in transactions[:25]:  # Show top 25 recent
        created_str = t.created_at.strftime("%d.%m %H:%M") if t.created_at else ""
        sign = "+" if t.type == "income" else "-"
        amt_style = cell_bold if t.type == "income" else cell_text
        tx_table_data.append([
            Paragraph(created_str, cell_text),
            Paragraph(type_labels.get(t.type, t.type), cell_text),
            Paragraph(t.category or "-", cell_text),
            Paragraph(f"{sign}{t.amount:,.0f}", amt_style),
            Paragraph(t.note or t.raw_expression or "-", cell_text),
        ])
        
    txt = Table(tx_table_data, colWidths=[70, 85, 95, 95, 177])
    txt.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
        ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(txt)
    
    doc.build(story)
    buffer.seek(0)
    return buffer
